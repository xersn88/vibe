from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "Monster_Cafe_详细实验报告.docx"
SCREENSHOT_DIR = ROOT / "report_assets" / "screenshots"


PROCESS_STEPS = [
    {
        "title": "项目骨架创建",
        "prompt": "从零创建 Monster Café 项目骨架，使用 Node.js + Express + SQLite，backend/frontend 分目录，后端托管前端页面。",
        "files": "新增 package.json、backend/server.js、backend/database.js、backend/routes/api.js、frontend/index.html、game.html、collection.html、shop.html、logs.html 等基础页面。",
        "architecture": "形成 backend + frontend 的基本全栈目录结构，Express 作为唯一入口，同时提供 API 和静态页面。",
        "work": "实现 npm install / npm start 的本地运行方式，完成服务启动、静态资源托管和基础页面跳转。",
        "logic": "server.js 负责初始化数据库后启动服务；frontend 页面通过 app.js 统一请求后端接口。"
    },
    {
        "title": "SQLite 数据库实现",
        "prompt": "网页交互文字中文化；实现 players、monsters、ingredients、orders、collection、dev_logs 等表，并插入默认怪兽和原料。",
        "files": "重点修改 backend/database.js，同时补充前端中文展示文案。",
        "architecture": "数据库从空壳变成项目数据中心，保存玩家、怪兽、原料、订单、图鉴和日志。",
        "work": "实现初始化函数，启动 server.js 时自动建表，并插入默认怪兽和默认原料数据。",
        "logic": "使用 CREATE TABLE IF NOT EXISTS 保证重复启动不报错；后续通过 ensureColumn 兼容字段扩展。"
    },
    {
        "title": "随机顾客与订单生成",
        "prompt": "新增 GET /api/game/next-customer，随机选择怪兽，根据喜欢口味生成有趣订单，支持 8 种口味标签。",
        "files": "修改 backend/routes/api.js，补充怪兽与口味标签相关数据。",
        "architecture": "游戏从静态页面进入后端驱动的随机订单阶段。",
        "work": "实现随机怪兽选择、订单文本拼接、推荐口味标签返回。",
        "logic": "后端读取 monsters，按怪兽喜欢口味生成需求文本，例如冰、甜、气泡、星光等标签组合。"
    },
    {
        "title": "饮品调配和评分算法",
        "prompt": "实现 POST /api/game/make-drink，提交 monsterId 和原料数组，按喜欢/讨厌口味评分，保存订单并更新玩家金币和图鉴。",
        "files": "修改 backend/routes/api.js，扩展 orders、collection 更新逻辑。",
        "architecture": "形成“顾客需求 - 玩家选料 - 后端评分 - 数据更新”的核心闭环。",
        "work": "实现满意度、金币奖励、订单记录、库存扣减、图鉴来店次数和最高满意度更新。",
        "logic": "后端将选中原料标签合并成集合，计算命中喜欢标签和讨厌标签，再生成评价文本。"
    },
    {
        "title": "主游戏页面",
        "prompt": "实现 game.html，显示金币、天数、等级、当前怪兽、订单需求、原料按钮、制作结果和下一位顾客。",
        "files": "修改 frontend/game.html、frontend/app.js、frontend/styles.css。",
        "architecture": "前端从基础页面变成可交互游戏界面。",
        "work": "绑定生成顾客、点击原料加入饮品、调用 make-drink、展示满意度和奖励。",
        "logic": "app.js 维护 currentCustomer、currentDrinkIngredients 等局部状态，按钮事件触发接口请求并刷新 UI。"
    },
    {
        "title": "怪兽图鉴",
        "prompt": "提供 GET /api/collection，collection.html 用卡片展示已解锁怪兽，未解锁显示问号卡片。",
        "files": "修改 backend/routes/api.js、frontend/collection.html、frontend/app.js、frontend/styles.css。",
        "architecture": "增加独立图鉴页面，游戏结果开始反向影响长期收集内容。",
        "work": "展示怪兽名称、种族、描述、来店次数、最高满意度。",
        "logic": "图鉴数据由 save_collection 或 collection 表驱动，未解锁时隐藏真实信息。"
    },
    {
        "title": "原料商店",
        "prompt": "实现 GET /api/ingredients 和 POST /api/shop/buy，支持购买库存和金币解锁原料。",
        "files": "修改 backend/routes/api.js、frontend/shop.html、frontend/app.js、frontend/styles.css。",
        "architecture": "加入经营资源循环：金币可以换库存，库存影响制作饮品。",
        "work": "商店页面展示原料名称、标签、价格、库存和是否解锁，购买后更新金币与库存。",
        "logic": "购买接口校验金币是否足够，再扣金币、增加库存或解锁状态。"
    },
    {
        "title": "修复开始游戏流程",
        "prompt": "基于已有阶段让游戏真正可运行，可以生成顾客、交互并完成游戏。",
        "files": "主要修改 frontend/app.js 与 backend/routes/api.js 中的启动和接口调用细节。",
        "architecture": "不改变总体架构，修复前后端衔接。",
        "work": "排查无法生成顾客的问题，修复页面初始化和接口请求顺序。",
        "logic": "页面加载时先取玩家和原料，再请求下一位顾客，避免空状态直接进入制作。"
    },
    {
        "title": "日结、更多顾客和成就",
        "prompt": "增加更多顾客，补齐原料标签覆盖；加入一天结束结算、星级评分、经验升级、成就和奖励。",
        "files": "扩展 backend/database.js 默认数据，修改 backend/routes/api.js、frontend/game.html、frontend/app.js、styles.css。",
        "architecture": "从单订单循环扩展为按天经营循环，增加长期成长系统。",
        "work": "新增更多怪兽和原料，加入日结弹窗、利润、顾客数、满意度星级、经验和成就。",
        "logic": "一天结束时汇总 dayStats，按利润、顾客数、平均满意度计算星级和奖励。"
    },
    {
        "title": "结算错误与评分修正",
        "prompt": "修复结算错误；调整评分机制，避免全覆盖标签仍然只有低分。",
        "files": "修改 backend/routes/api.js 的结算和评分相关逻辑。",
        "architecture": "不改变页面结构，优化核心规则。",
        "work": "修复资源查找错误，调整满意度计算，使覆盖需求标签的饮品能得到合理高分。",
        "logic": "评分从简单加减分调整为需求覆盖、喜欢标签、讨厌标签和组合奖励共同作用。"
    },
    {
        "title": "多标签平等与可解订单",
        "prompt": "第一、第二标签地位应均等；顾客需求必须存在满分可能，避免无解。",
        "files": "修改 backend/routes/api.js 的标签解析、订单生成和评分函数。",
        "architecture": "游戏规则层变得更严格，订单生成前增加可解性判断。",
        "work": "统一解析原料多标签，生成订单时基于已解锁原料标签判断是否能避开讨厌口味并覆盖需求。",
        "logic": "把原料标签统一拆分为数组/集合，任意位置命中都算有效；订单只从可满足组合中抽取。"
    },
    {
        "title": "暂停、补货、等待和差评",
        "prompt": "加入暂停弹窗、补货状态、补货返回、缺货标识、顾客等待条、超时离开和差评。",
        "files": "修改 frontend/game.html、shop.html、app.js、styles.css 和 backend/routes/api.js。",
        "architecture": "加入局内临时状态保存，商店可以作为营业中的补货页面使用。",
        "work": "实现暂停/继续/退出本天/前往商店补货，缺货按钮置灰，顾客等待超时产生差评。",
        "logic": "补货前保存 gameState 到 localStorage，商店识别 restock 状态隐藏导航，返回后恢复计时和当前顾客。"
    },
    {
        "title": "读取新要求文档",
        "prompt": "读取包含新要求的文本文档并按要求更改。",
        "files": "读取 更新要求.txt，并据此继续修改页面流程和功能模块。",
        "architecture": "需求方向从单一游戏页扩展为完整应用流程。",
        "work": "根据文档内容重构首页、存档、新手教程和页面导航思路。",
        "logic": "把需求拆分为标题页、存档页、主页、功能页和局内页面几个层级。"
    },
    {
        "title": "标题页、存档页和导航重构",
        "prompt": "新增初始标题页和存档页；index 改为读档后的主页；图鉴拆成怪兽图鉴和成就图鉴；游戏页重新布局。",
        "files": "新增 frontend/start.html、saves.html、achievements.html；重写 index.html、game.html、app.js、styles.css。",
        "architecture": "页面流转变为 start.html → saves.html → index.html → game.html，各功能页挂在存档主页导航下。",
        "work": "实现标题页艺术字、存档读取/删除/新建、主页玩家信息、导航拆分和游戏页上中下布局。",
        "logic": "localStorage 保存当前存档 ID，所有接口请求通过 X-Save-Id 传给后端。"
    },
    {
        "title": "多存档隔离与背包",
        "prompt": "修复存档互通；每个新档独立金币、经验、图鉴、库存、成就；新增背包页面和奖励弹窗。",
        "files": "修改 database.js 新增 save_* 表，新增 backpack.html，修改 app.js、api.js、styles.css。",
        "architecture": "从全局玩家数据改为以 save_id 为核心的数据隔离架构。",
        "work": "实现 save_ingredients、save_collection、save_achievements、save_backpack、save_level_rewards 等存档维度表。",
        "logic": "后端每个接口先读取 X-Save-Id，并通过 ensureSaveScopedData 初始化当前存档的数据。"
    },
    {
        "title": "教程不计入正式经营",
        "prompt": "新手教程不能算营业第一天，也不能解锁怪兽或计入来店次数；成就待办排序和日结入账要修正。",
        "files": "修改 backend/routes/api.js、database.js、frontend/app.js。",
        "architecture": "增加“教程模式”和“正式营业模式”的状态区分。",
        "work": "新档主页显示已营业 0 天，怪兽全未解锁；成就按可领取、未完成、已完成排序。",
        "logic": "教程流程只演示操作，不写入正式收益、经验、好感和图鉴；正式收益在日结时统一提交。"
    },
    {
        "title": "空存档、返回标题和加载动画",
        "prompt": "允许删到无存档；index 的主页按钮改为返回标题；外部页面加载时显示布丁加载动画。",
        "files": "修改 saves.html、index.html、app.js、styles.css。",
        "architecture": "存档系统支持空列表状态；页面切换体验增加全局加载层。",
        "work": "删除最后一个存档不再强制保留；返回标题清理登录状态；加载时显示布丁图标、进度条和随机提示语。",
        "logic": "加载动画只在非营业页面触发，game.html 营业中不显示，避免影响计时和交互。"
    },
    {
        "title": "重复问题复核",
        "prompt": "再次确认新档待办、删除最后存档、返回标题和加载动画问题。",
        "files": "继续核对 app.js 和相关页面。",
        "architecture": "无新架构变化，属于稳定性修复。",
        "work": "复核红点逻辑、删除逻辑和页面加载流程。",
        "logic": "待办只对“已完成未领取”或“已解锁未读”内容提示，不对未完成内容提示。"
    },
    {
        "title": "新手教程完整流程",
        "prompt": "新建存档后先播放营业过场，再弹新手教程选择；设计模块式教程和教程模拟营业。",
        "files": "新增/修改 intro.html、index.html、game.html、app.js、styles.css。",
        "architecture": "加入 intro 过场和教程引导层。",
        "work": "实现教程选择弹窗、模块高亮对话框、必须点击高亮模块推进、教程模拟营业和初始奖励弹窗。",
        "logic": "说明型步骤显示下一步按钮，交互型步骤拦截其他点击，只允许点击指定模块。"
    },
    {
        "title": "进入下一天 bug 修复",
        "prompt": "完成前一天点击进入下一天后无法开始，天数也不更新。",
        "files": "修改 frontend/app.js 和后端日结返回数据。",
        "architecture": "无结构变化，修复日结到下一天的状态机。",
        "work": "让结算弹窗的“进入下一天”同时清理旧日状态并启动新一天。",
        "logic": "日结接口更新 business_days，前端 next-day 按 business_days + 1 显示局内第几天。"
    },
    {
        "title": "进入下一天问题复测",
        "prompt": "重复提交进入下一天 bug，要求再次修复无法开始和天数不更新。",
        "files": "复核 app.js 的 startNextDay、beginNewDay 和玩家状态加载。",
        "architecture": "无结构变化，属于回归验证。",
        "work": "确认日结后天数更新、进入下一天自动开始、局内第几天等于已营业天数 + 1。",
        "logic": "区分主页“已营业几天”和局内“第几天”。"
    },
    {
        "title": "特殊道具 SQL bug",
        "prompt": "根据截图修复特殊道具制作饮品时报 SQLITE no such column 的问题，并检查类似错误。",
        "files": "修改 backend/routes/api.js。",
        "architecture": "无结构变化，修正 SQL 更新条件。",
        "work": "修复把 save_collection 更新条件误写成 monster_affinity.save_id 的问题。",
        "logic": "所有按存档更新的 SQL 都必须使用当前表真实字段和 save_id 条件。"
    },
    {
        "title": "收尾、删除无用页面和日志页",
        "prompt": "删除设置页和没用上的页面，全面检查网站；实现 AI 协作日志页面的增删查。",
        "files": "删除 settings.html、report.html；修改 logs.html、app.js、api.js、database.js、README.md。",
        "architecture": "页面结构收敛为实际使用页面；开发日志成为正式功能模块。",
        "work": "实现 /api/logs GET/POST/DELETE，logs.html 可新增、展示、删除开发日志。",
        "logic": "dev_logs 保存提示词、AI 返回摘要、人工修改和运行结果，用于课程展示开发过程。"
    },
    {
        "title": "整理全部提示词入日志",
        "prompt": "整理本对话中全部提示词，精简后加入所有日志，不遗漏任何一次对话。",
        "files": "新增 scripts/seed-dev-logs.js 并写入 dev_logs 数据。",
        "architecture": "无页面架构变化，补充演示数据。",
        "work": "整理 23 条阶段日志，覆盖从项目骨架到收尾日志页的完整开发过程。",
        "logic": "脚本先清空 dev_logs，再按实际表结构兼容旧列并批量写入 UTF-8 中文内容。"
    },
    {
        "title": "修复日志乱码和时间显示",
        "prompt": "开发日志全是乱码，并且不想标出时间。",
        "files": "修改 frontend/app.js、backend/routes/api.js、scripts/seed-dev-logs.js。",
        "architecture": "无架构变化，修复展示和数据写入。",
        "work": "重新用 UTF-8 写入日志，日志接口不再返回 created_at，前端卡片不再显示时间。",
        "logic": "乱码原因是之前批量导入时编码被破坏；修复后直接用 Node/SQLite 写入 Unicode 字符串。"
    },
    {
        "title": "实验报告初稿",
        "prompt": "创建详细实验报告，包含任务描述、实验过程和收获心得，图文并茂并插入运行截图。",
        "files": "新增 scripts/capture-report-screenshots.js、scripts/build-experiment-report.py，生成 Monster_Cafe_实验报告.docx。",
        "architecture": "项目代码不变，增加报告生成辅助脚本和 report_assets 截图目录。",
        "work": "使用 Playwright/Edge 运行本地项目截图，并用 python-docx 生成 Word 文档。",
        "logic": "脚本准备报告演示存档，自动访问各页面截图，再将截图插入报告。"
    },
    {
        "title": "实验报告详细化",
        "prompt": "报告开头说明智能体辅助；实验过程写明每次提示词、每步工作、增删改文件、架构变化、代码逻辑；补齐教程、弹窗、重要代码附录。",
        "files": "扩展截图脚本和报告生成脚本，重新生成 Monster_Cafe_实验报告.docx。",
        "architecture": "项目代码不变，报告内容从概览版升级为按对话顺序展开的详细版。",
        "work": "补充新手教程、加载过场、返回标题、退出确认、日结、奖励弹窗、图鉴详情等截图，并增加代码附录。",
        "logic": "报告先讲 Codex 智能体如何辅助，再按提示词顺序复盘开发，最后展示最终运行效果。"
    },
]


SCREENSHOTS = [
    ("01_start.png", "图 1 标题页：艺术字游戏标题与开始/退出入口"),
    ("02_exit_game_modal.png", "图 2 退出游戏确认弹窗"),
    ("02_saves.png", "图 3 存档页：读取、新建和删除存档入口"),
    ("03_new_save_modal.png", "图 4 新建存档弹窗：输入昵称或随机昵称"),
    ("04_delete_save_modal.png", "图 5 删除存档确认弹窗"),
    ("03_home.png", "图 6 存档主页：玩家信息、等级经验和成就收集进度"),
    ("05_loading_overlay.png", "图 7 外部页面加载过场：布丁图标、进度条和提示语"),
    ("06_tutorial_choice_modal.png", "图 8 新手教程选择弹窗"),
    ("07_home_tutorial_guide.png", "图 9 首页模块式新手引导"),
    ("08_return_title_modal.png", "图 10 返回标题确认弹窗"),
    ("09_game_countdown.png", "图 11 开始营业倒计时"),
    ("04_game_customer.png", "图 12 游戏页：顾客、订单、等待条、制作台、原料区"),
    ("10_game_tutorial_guide.png", "图 13 教程营业中的模块引导"),
    ("05_game_result.png", "图 14 制作饮品后的评分、奖励和评价墙"),
    ("06_pause.png", "图 15 暂停弹窗"),
    ("11_exit_day_confirm.png", "图 16 退出本天确认弹窗"),
    ("12_day_settlement_modal.png", "图 17 今日结算弹窗：星级、收入、成本、经验和成就"),
    ("07_shop.png", "图 18 原料商店"),
    ("08_restock_shop.png", "图 19 补货状态商店"),
    ("09_collection.png", "图 20 怪兽图鉴总览"),
    ("15_collection_detail.png", "图 21 怪兽图鉴详情与故事列表"),
    ("16_story_reward_popup.png", "图 22 阅读故事后的奖励弹窗"),
    ("10_achievements.png", "图 23 成就图鉴"),
    ("17_achievement_reward_popup.png", "图 24 领取成就奖励弹窗"),
    ("11_level_rewards.png", "图 25 等级奖励页面"),
    ("18_level_reward_popup.png", "图 26 领取等级奖励弹窗"),
    ("12_backpack.png", "图 27 背包页面"),
    ("13_logs_form.png", "图 28 AI 协作日志新增表单"),
    ("14_logs_list.png", "图 29 AI 协作日志列表"),
]


CODE_SNIPPETS = [
    (
        "Express 服务入口与静态托管",
        "backend/server.js",
        dedent(
            """
            app.get('/', (req, res) => {
              res.sendFile(path.join(frontendDir, 'start.html'));
            });

            app.use('/api', apiRoutes);
            app.use(express.static(frontendDir, {
              etag: false,
              lastModified: false,
              maxAge: 0
            }));
            """
        ).strip(),
    ),
    (
        "前端请求自动携带当前存档 ID",
        "frontend/app.js",
        dedent(
            """
            async function fetchJson(url, options = {}) {
              const selectedSave = localStorage.getItem('monsterCafeSelectedSave');
              const headers = { ...(options.headers || {}) };
              if (selectedSave) {
                headers['X-Save-Id'] = selectedSave;
              }
              const response = await fetch(url, { ...options, headers });
              return response.json();
            }
            """
        ).strip(),
    ),
    (
        "存档维度数据初始化",
        "backend/routes/api.js",
        dedent(
            """
            async function ensureSaveScopedData(saveId) {
              const save = await get('SELECT id FROM saves WHERE id = ?', [saveId]);
              if (!save) return false;
              const ingredients = await all('SELECT id, unlocked FROM ingredients');
              for (const ingredient of ingredients) {
                await run(
                  'INSERT OR IGNORE INTO save_ingredients (save_id, ingredient_id, stock, unlocked) VALUES (?, ?, ?, ?)',
                  [saveId, ingredient.id, ingredient.unlocked ? 2 : 0, ingredient.unlocked ? 1 : 0]
                );
              }
              return true;
            }
            """
        ).strip(),
    ),
    (
        "随机顾客与订单生成接口",
        "backend/routes/api.js",
        dedent(
            """
            router.get('/game/next-customer', async (req, res, next) => {
              const saveId = getRequestSaveId(req);
              await ensureSaveScopedData(saveId);
              const monsters = await all('SELECT * FROM monsters WHERE unlock_level <= ? ORDER BY RANDOM()', [playerLevel]);
              const monster = pickWeightedMonster(monsters);
              const order = buildSolvableOrder(monster, unlockedIngredients);
              res.json({ monster, orderText: order.text, requestedTasteTags: order.tags });
            });
            """
        ).strip(),
    ),
    (
        "饮品评分核心思路",
        "backend/routes/api.js",
        dedent(
            """
            const selectedTags = new Set();
            for (const ingredient of selectedIngredients) {
              parseTags(ingredient.taste_tags).forEach((tag) => selectedTags.add(tag));
            }
            const hitRequired = requiredTags.filter((tag) => selectedTags.has(tag)).length;
            const dislikedHits = dislikedTags.filter((tag) => selectedTags.has(tag)).length;
            let satisfaction = 40 + hitRequired * 25 - dislikedHits * 20 + comboBonus;
            satisfaction = Math.max(0, Math.min(100, satisfaction));
            """
        ).strip(),
    ),
    (
        "日结入账与成长",
        "backend/routes/api.js",
        dedent(
            """
            const profitStars = getStars(profit, [100, 70, 40, 15]);
            const customerStars = getStars(reviewedCustomers, [8, 6, 4, 2]);
            const satisfactionStars = getStars(avgSatisfaction, [90, 80, 65, 50]);
            const totalXpGained = drinkXpTotal + settlementXp;
            await run(
              'UPDATE saves SET business_days = business_days + 1, coins = coins + ?, xp = xp + ?, level = ? WHERE id = ?',
              [totalCoinGained, totalXpGained, nextLevel, saveId]
            );
            """
        ).strip(),
    ),
    (
        "开发日志接口不返回时间",
        "backend/routes/api.js",
        dedent(
            """
            router.get('/logs', async (req, res, next) => {
              const logs = await all(`
                SELECT id, prompt, ai_summary, manual_change, run_result
                FROM dev_logs
                ORDER BY id ASC
              `);
              res.json(logs);
            });
            """
        ).strip(),
    ),
]


def set_east_asia_font(run, font_name="宋体"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")


def add_paragraph(doc, text="", size=12, bold=False, color=None, align=None, before=0, after=6):
    paragraph = doc.add_paragraph()
    if text:
        run = paragraph.add_run(text)
        set_east_asia_font(run)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_east_asia_font(run, "黑体")
    run.bold = True
    run.font.color.rgb = RGBColor(46, 33, 67)
    run.font.size = Pt({1: 16, 2: 14, 3: 12.5}.get(level, 12.5))
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=10.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], "FCE8F1")
        set_cell_text(table.rows[0].cells[index], header, bold=True, size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=font_size)
    add_paragraph(doc, "", after=2)
    return table


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_east_asia_font(run)
        run.font.size = Pt(12)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_east_asia_font(run)
        run.font.size = Pt(12)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.space_after = Pt(3)


def add_image(doc, filename, caption, width=6.1):
    image_path = SCREENSHOT_DIR / filename
    if not image_path.exists():
        add_paragraph(doc, f"（截图缺失：{filename}）", color="9B1C1C")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    caption_p = add_paragraph(doc, caption, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color="6B4C72", after=8)
    caption_p.paragraph_format.keep_with_next = True


def add_code_block(doc, title, path, code):
    add_heading(doc, title, 3)
    add_paragraph(doc, f"来源文件：{path}", size=10.5, color="6B4C72", after=3)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(8)
    for line in code.splitlines():
      run = paragraph.add_run(line + "\n")
      run.font.name = "Consolas"
      run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
      run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
      run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
      run.font.size = Pt(8.5)
      run.font.color.rgb = RGBColor(31, 31, 31)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Monster Café 怪兽咖啡馆经营小游戏详细实验报告")
    set_east_asia_font(run, "黑体")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor(46, 33, 67)
    title.paragraph_format.space_after = Pt(12)

    add_paragraph(
        doc,
        "基于 HTML、CSS、JavaScript、Node.js、Express 与 SQLite 的本地网页小游戏系统",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color="6B4C72",
        after=14,
    )

    add_table(
        doc,
        ["项目项", "内容"],
        [
            ["项目名称", "Monster Café 怪兽咖啡馆经营小游戏"],
            ["开发方式", "用户提出阶段性需求，Codex 智能体协助完成需求拆解、代码实现、调试、截图和报告整理"],
            ["技术栈", "前端 HTML/CSS/JavaScript；后端 Node.js + Express；数据库 SQLite"],
            ["运行方式", "npm install 安装依赖，npm start 启动本地服务"],
            ["报告范围", "按对话顺序复盘全部主要提示词、工作内容、文件变化、架构变化、机制说明、最终截图和关键代码"],
        ],
    )

    add_heading(doc, "一、智能体辅助说明", 1)
    add_paragraph(
        doc,
        "本项目开发过程中使用 Codex 作为全栈开发智能体辅助。我的角色是提出课程设计目标、功能需求、界面反馈和测试问题；Codex 的角色是根据需求阅读项目文件、设计数据结构、编写前后端代码、修复运行问题、补充测试数据、运行本地服务截图，并最终协助整理实验报告。"
    )
    add_bullets(
        doc,
        [
            "需求拆解：把“怪兽咖啡馆经营小游戏”拆分为项目骨架、数据库、游戏接口、主页面、图鉴、商店、成就、背包、日志和报告等阶段。",
            "代码实现：按阶段修改 Express 路由、SQLite 初始化逻辑、HTML 页面、CSS 样式和原生 JavaScript 交互。",
            "调试修复：根据截图和运行现象定位结算错误、评分不合理、多存档数据互通、补货返回失效、SQL 字段错误、日志乱码等问题。",
            "展示整理：使用本地服务和浏览器自动化生成页面截图，并将实验过程、功能结果和关键代码整理到 Word 报告中。",
        ],
    )

    add_heading(doc, "二、任务描述", 1)
    add_paragraph(
        doc,
        "Monster Café 是一个怪兽咖啡馆经营小游戏。玩家扮演咖啡馆老板，每天接待不同小怪兽顾客。顾客会提出饮品口味需求，玩家选择不同原料调配饮品，系统根据原料标签和顾客偏好计算满意度、金币奖励、经验和长期收集进度。项目还包括库存管理、怪兽图鉴、原料商店、订单记录、成就图鉴、等级奖励、背包和 AI 协作开发日志。"
    )
    add_bullets(
        doc,
        [
            "前端使用 HTML、CSS、JavaScript，不使用复杂框架。",
            "后端使用 Node.js + Express，数据库使用 SQLite。",
            "项目能够本地运行，后端托管 frontend 静态页面。",
            "界面风格偏可爱、像素风、卡片式布局，交互文字为中文，特殊名称如 Monster Café 和怪兽名保留。",
            "代码结构清晰，适合课程设计展示和答辩。",
        ],
    )

    add_heading(doc, "三、最初代码架构", 1)
    add_paragraph(
        doc,
        "项目最初按 backend 和 frontend 两个主目录组织。backend 内包含 server.js、database.js 和 routes 文件夹；frontend 内包含各个 HTML 页面、统一样式 styles.css 和统一交互脚本 app.js。这个架构的优点是简单直接，适合课程设计展示：后端负责数据和接口，前端负责页面与交互。"
    )
    add_table(
        doc,
        ["目录或文件", "初始职责", "后续变化"],
        [
            ["backend/server.js", "Express 服务入口，托管静态页面", "保持入口职责，继续挂载 /api 路由和 frontend 静态资源"],
            ["backend/database.js", "SQLite 初始化", "逐步扩展多存档、成就、背包、等级奖励、好感、礼物等表"],
            ["backend/routes/api.js", "基础 API 路由", "成为游戏逻辑核心，承载顾客生成、评分、商店、图鉴、成就、日志等接口"],
            ["frontend/*.html", "首页和功能页面", "扩展为标题页、存档页、主页、游戏页、商店、图鉴、成就、背包、日志等完整页面"],
            ["frontend/app.js", "前端交互脚本", "承担页面初始化、局内状态、教程、弹窗、计时、待办红点和接口调用"],
            ["frontend/styles.css", "基础样式", "扩展为像素风卡片、弹窗、加载层、教程高亮、奖励弹窗等完整视觉系统"],
        ],
    )

    add_heading(doc, "四、按对话顺序的实验过程", 1)
    add_paragraph(
        doc,
        "本节按对话中每一次主要提示词的顺序展开。每个步骤都记录了精简版提示词、完成的工作、增删改文件、架构变化和主要实现机制。"
    )

    for index, step in enumerate(PROCESS_STEPS, start=1):
        add_heading(doc, f"{index}. {step['title']}", 2)
        add_paragraph(doc, f"精简提示词：{step['prompt']}")
        add_paragraph(doc, f"完成工作：{step['work']}")
        add_paragraph(doc, f"增删改文件：{step['files']}")
        add_paragraph(doc, f"架构变化：{step['architecture']}")
        add_paragraph(doc, f"逻辑机制：{step['logic']}")

    add_heading(doc, "五、最终运行效果展示", 1)
    add_paragraph(
        doc,
        "完成所有开发与修复后，系统形成了完整的本地可运行网页小游戏。以下截图均通过本地 Express 服务运行页面后截取，覆盖标题、存档、主页、教程、游戏、弹窗、结算、商店、图鉴、成就、等级奖励、背包和开发日志等主要功能。"
    )
    for filename, caption in SCREENSHOTS:
        add_image(doc, filename, caption, width=5.95)

    add_heading(doc, "六、接口与功能验证", 1)
    add_table(
        doc,
        ["模块", "主要接口或页面", "验证结果"],
        [
            ["存档系统", "GET/POST/DELETE /api/saves，saves.html", "支持新建、读取、删除到空存档；当前存档 ID 经 X-Save-Id 传递"],
            ["游戏流程", "/api/game/next-customer，/api/game/make-drink，game.html", "可生成顾客、选择原料、制作饮品、显示评分和评价"],
            ["评分算法", "make-drink 内部评分逻辑", "多标签平等，需求可解，命中需求加分，命中讨厌口味扣分"],
            ["日结系统", "/api/game/end-day", "可计算收入、成本、利润、平均满意度、星级、经验和成就"],
            ["商店补货", "/api/ingredients，/api/shop/buy，shop.html", "可购买库存、解锁原料；补货状态会暂停营业并隐藏普通导航"],
            ["怪兽图鉴", "/api/collection，/api/collection/todos", "支持未解锁问号卡、已解锁详情、故事阅读和奖励弹窗"],
            ["成就图鉴", "/api/achievements，/api/achievements/:id/claim", "完成未领取前置排序，领取后弹出奖励并消除待办"],
            ["等级奖励", "/api/level-rewards，/api/level-rewards/:id/claim", "按等级展示，可领取奖励显示待办，已领取项目后置"],
            ["背包", "/api/backpack，backpack.html", "展示原料、礼物和特殊物品，点击后显示详情"],
            ["开发日志", "/api/logs，logs.html", "支持新增、查询、删除；中文正常，不显示时间"],
        ],
    )

    add_heading(doc, "七、收获与心得", 1)
    add_paragraph(
        doc,
        "通过本实验，我对一个小型全栈网页项目从需求到实现的完整流程有了更清晰的认识。最初的需求只是一个怪兽咖啡馆经营小游戏，但随着测试和反馈逐步增加，项目扩展出了多存档、教程、待办、奖励、背包和开发日志等系统。这个过程说明课程设计不仅要能运行，还要能解释清楚数据从哪里来、状态如何变化、用户操作如何影响后端数据。"
    )
    add_bullets(
        doc,
        [
            "需求层面：阶段性提示词能让复杂系统逐步成形，每次反馈都推动了规则、界面或数据结构变得更合理。",
            "架构层面：多存档隔离是一次重要重构，必须把金币、经验、库存、图鉴、成就、背包和等级奖励都放到 save_id 维度下。",
            "算法层面：评分机制要符合玩家直觉，不能只看第一个标签，也不能生成无法满分的订单。",
            "交互层面：暂停、补货、等待条、弹窗、红点待办和加载过场让游戏从功能可用变成体验完整。",
            "调试层面：截图反馈非常重要，结算错误、SQL 字段错误、补货返回失效和日志乱码都是通过运行现象定位并修复的。",
            "AI 协作层面：Codex 能提高实现和排查效率，但需求判断、结果验收和最终取舍仍然需要由开发者主导。",
        ],
    )

    add_heading(doc, "八、本地运行方式", 1)
    add_numbered(
        doc,
        [
            "进入项目目录：C:\\Users\\xersn\\Documents\\vibe。",
            "执行 npm install 安装依赖。",
            "执行 npm start 启动 Express 服务。",
            "浏览器访问 http://localhost:3000，进入 Monster Café 标题页。",
            "选择或新建存档后进入主页，可以开始营业、补货、查看图鉴、领取奖励、查看背包和记录开发日志。",
        ],
    )

    add_heading(doc, "附录：关键代码摘录", 1)
    add_paragraph(
        doc,
        "以下代码为报告说明用的关键片段，展示项目的主要机制。为了便于阅读，附录中只截取核心逻辑，省略了部分异常处理和展示细节。"
    )
    for title, path, code in CODE_SNIPPETS:
        add_code_block(doc, title, path, code)

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_document())
