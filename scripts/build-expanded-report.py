from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BASE_SCRIPT = ROOT / "scripts" / "build-experiment-report.py"
OUT_PATH = ROOT / "Monster_Cafe_扩写实验报告.docx"
SCREENSHOT_DIR = ROOT / "report_assets" / "screenshots"

spec = spec_from_file_location("base_report", BASE_SCRIPT)
base_report = module_from_spec(spec)
spec.loader.exec_module(base_report)
PROCESS_STEPS = base_report.PROCESS_STEPS


DETAIL_NOTES = {
    "项目骨架创建": [
        "这一阶段的重点不是先堆功能，而是先把项目运行链路搭稳：Node 负责启动服务，Express 负责路由和静态资源，frontend 只放页面和静态脚本。这样后续每次增加功能，都能明确知道前端文件和后端文件分别该放在哪里。",
        "代码实现时先建立最小可运行闭环：package.json 声明依赖和 start 命令；server.js 启动 Express；database.js 预留初始化函数；routes/api.js 作为 API 的集中入口；前端页面先按要求列全，避免后面频繁调整目录。",
        "这个阶段的架构变化最大，因为它从无项目变成了可运行工程。后续所有功能都建立在这个目录约定上：后端只处理数据和接口，前端只处理显示和交互，二者通过 /api 通信。"
    ],
    "SQLite 数据库实现": [
        "数据库设计的核心思路是先覆盖课程要求的基础实体，再给后续功能留下扩展空间。players 保存玩家经营状态，monsters 保存顾客设定，ingredients 保存原料，orders 保存调饮结果，collection 保存图鉴，dev_logs 保存开发过程。",
        "初始化逻辑采用“可重复执行”的写法。CREATE TABLE IF NOT EXISTS 确保反复 npm start 不会重建表；默认数据插入前先检查是否存在，避免重复插入同名怪兽和原料。",
        "后续项目多次加字段，因此 database.js 里保留 ensureColumn 这种迁移式逻辑。它让数据库在已有旧数据时也能补齐新字段，适合课程项目迭代频繁的情况。"
    ],
    "随机顾客与订单生成": [
        "顾客生成不是简单随机名字，而是要把怪兽设定和订单需求绑定起来。后端先随机抽取怪兽，再读取它喜欢和讨厌的口味，最后生成符合角色口吻的订单文本。",
        "订单文本采用模板化生成，例如把“冰”“甜”“气泡”等标签转成“冰冰的、甜甜的、会冒泡的饮品”。这样既保证订单可被程序计算，又让玩家看到的是自然的中文需求。",
        "推荐标签会和订单文本一起返回给前端。前端显示这些标签，玩家据此选择原料，后端仍然作为最终评分者，避免前端直接决定结果。"
    ],
    "饮品调配和评分算法": [
        "这个接口是游戏闭环的核心。前端只提交 monsterId 和 selectedIngredients，后端重新查询怪兽和原料，保证评分使用数据库中的真实数据，而不是信任前端传来的标签。",
        "评分时把所有选中原料的标签合并成集合，再分别计算需求命中、喜欢口味命中和讨厌口味命中。命中喜欢或需求提高满意度，命中讨厌口味降低满意度，最后把分数限制在合理范围。",
        "制作成功后不仅返回结果，还要更新多个长期数据：扣库存、写订单、更新金币、更新图鉴来店次数和最高满意度。这保证一次调饮能影响后续经营。"
    ],
    "主游戏页面": [
        "主游戏页的工作思路是把玩家最常用的信息放在一屏内：顶部显示经营状态，左侧显示顾客，中间显示订单和制作台，右侧显示评价，下方横向显示原料。",
        "前端使用原生 JavaScript 维护当前局内状态，例如当前顾客、当前订单文本、当前已选原料、当前礼物、今日统计等。每次点击原料按钮只改变局部状态，点击制作饮品才提交后端。",
        "页面没有使用框架，因此事件绑定和渲染函数要分清职责。loadPlayerStatus 负责状态栏，renderCustomer 负责顾客卡，renderIngredientButtons 负责原料区，makeDrink 负责提交制作。"
    ],
    "怪兽图鉴": [
        "图鉴功能的关键是区分“数据库里存在的怪兽”和“当前存档已解锁的怪兽”。未解锁怪兽不能直接暴露信息，因此前端根据 unlocked 字段显示问号卡片。",
        "后端 collection 接口把 monsters、save_collection、monster_affinity 等数据组合返回，前端只负责展示。这样来店次数、最高满意度、好感度和故事状态都能集中在同一张卡片里。",
        "图鉴后来又扩展了故事待办逻辑：不是所有未完成内容都提示，而是已解锁、已达到条件、但未读的故事才显示红点。这让待办提示更准确。"
    ],
    "原料商店": [
        "商店是资源循环的一部分。玩家通过营业获得金币，再用金币补充库存或解锁新原料；新原料带来更多标签组合，又反过来提高完成订单的能力。",
        "购买接口需要校验三件事：原料是否存在、当前存档是否已解锁、金币是否足够。校验通过后再扣金币、加库存或改 unlocked 状态。",
        "前端商店卡片显示标签、价格、库存和状态。未解锁原料显示解锁按钮，已解锁原料显示购买按钮，避免玩家误操作。"
    ],
    "修复开始游戏流程": [
        "最初问题是页面能打开，但游戏无法真正生成顾客。这类问题通常不是单个函数错误，而是页面初始化顺序、接口路径和状态判断共同造成的。",
        "排查思路是先确认后端 next-customer 接口能返回数据，再确认前端是否在正确时机调用接口，最后看渲染函数是否收到 currentCustomer。",
        "修复后，页面加载先读取玩家和原料，再进入顾客生成。这样顾客出现时，原料按钮、订单文本和制作按钮都处于可用状态。"
    ],
    "日结、更多顾客和成就": [
        "这一阶段把游戏从“做一杯饮品”扩展成“经营一天”。因此需要引入今日统计 dayStats，包括顾客数、收入、成本、满意度、差评和经验。",
        "日结评分不只看金币收入，还要综合利润、顾客数量和平均满意度。这样玩家不能只追求高价，也要考虑服务效率和顾客体验。",
        "成就系统通过目标码记录进度，例如 first_day、five_customers、perfect_drink。后端在制作饮品或日结时更新进度，前端只展示可领取状态。"
    ],
    "结算错误与评分修正": [
        "结算错误来自资源查询和统计字段不一致。修复时先把前端 dayStats 发送的数据和后端 end-day 接收的数据对齐，再确认后端返回 summary 时字段完整。",
        "评分机制的调整重点是公平性。玩家如果已经覆盖订单需求，就不应该仍然只有很低分；因此基础分、需求命中、组合奖励和讨厌标签扣分需要重新平衡。",
        "新的评分逻辑让高分路径更清晰：覆盖需求标签、避开讨厌标签、控制原料数量，玩家就能稳定获得较高满意度。"
    ],
    "多标签平等与可解订单": [
        "用户指出第二标签也应该算命中，这是非常关键的规则修正。实现上不能只取 taste_tags 字符串的第一项，而要把所有标签 split 后放入集合。",
        "可解订单判断要基于已解锁原料，而不是当前库存。原因是库存不足时玩家可以补货，如果订单生成只看库存，商店补货机制就失去意义。",
        "生成订单前会检查需求标签是否能用已解锁原料覆盖，并尽量避免必须使用讨厌标签才能完成的组合，从源头减少无解订单。"
    ],
    "暂停、补货、等待和差评": [
        "暂停系统的核心是区分“现实页面跳转”和“游戏时间”。玩家进入商店补货时，网页确实离开 game.html，但游戏内计时必须暂停并能恢复。",
        "实现方式是把局内状态保存到 localStorage，包括剩余时间、顾客等待时间、当前顾客、订单文本、已选原料和今日统计。商店补货完成后再跳回 game.html 读取状态。",
        "顾客等待条增加了时间压力：等待超过一定时间心情下降，超时未制作会离开并产生差评。差评进入日结统计，降低满意度评价和经验收益。"
    ],
    "读取新要求文档": [
        "用户通过文本文档补充需求后，开发重点从单页游戏变成完整产品流程。我的处理方式是先读取文档，提炼出页面入口、存档、教程、导航和局内布局几条主线。",
        "这一阶段没有立即大改所有代码，而是先确定页面层级：标题页负责开始与退出，存档页负责选择身份，index 负责存档内主页，game 负责营业。",
        "这样拆分后，后续每个功能都有稳定入口，避免所有按钮都挤在一个首页里。"
    ],
    "标题页、存档页和导航重构": [
        "标题页 start.html 是玩家第一次进入游戏看到的页面，只保留开始游戏和退出游戏。开始游戏进入 saves.html，符合用户要求的“先选档再进入主页”。",
        "存档页支持单击读取、悬浮删除、新建存档和教程选择。index.html 不再是初始页面，而是读取存档后的咖啡馆柜台主页。",
        "游戏页布局也随之调整：上方放状态和三大功能区，下方放横向原料区。这个布局更适合调饮操作，因为玩家视线先看顾客和订单，再向下选择原料。"
    ],
    "多存档隔离与背包": [
        "多存档隔离是后期最重要的架构调整。原本很多数据类似全局数据，如果不改，A 存档获得的成就、金币或图鉴会影响 B 存档。",
        "解决方式是在关键数据表前加 save_ 前缀或增加 save_id 字段，例如 save_ingredients、save_collection、save_achievements、save_backpack、save_level_rewards。",
        "背包系统也按存档隔离。背包不仅显示原料，还显示礼物、故事奖励和特殊道具，点击物品后展示数量、简介和来源。"
    ],
    "教程不计入正式经营": [
        "教程模式和正式模式必须分开，否则新建存档会莫名其妙出现营业天数、怪兽来店次数和成就进度。修复时增加 tutorial 标记，让教程流程只演示操作。",
        "教程完成或跳过都会发初始奖励，但不增加 business_days，不解锁怪兽，不写正式订单收益，也不提升好感。",
        "日结入账也做了区分：局内可以显示临时金币和经验变化，但正式金币、经验、等级和成就要在一天结束后统一写入数据库。"
    ],
    "空存档、返回标题和加载动画": [
        "允许删到无存档是为了符合真实存档管理逻辑。删除最后一个存档后，页面应展示空状态和新建按钮，而不是强制保留一个默认档。",
        "index 页面上的“主页”按钮改为“返回标题”，因为当前已经在主页，再显示主页按钮没有意义。返回标题时清理 localStorage 中的登录态和工作进度。",
        "加载动画只在外部页面切换时出现，不在营业中出现。这样既能缓解页面切换等待感，又不会干扰游戏计时和顾客等待。"
    ],
    "重复问题复核": [
        "这次复核主要针对待办红点和新档初始状态。判断待办时不能把所有未完成成就都标红，否则玩家一进新档就会看到错误提示。",
        "正确逻辑是：成就只有 completed=1 且 claimed=0 才显示待办；怪兽故事只有已解锁、达到好感阈值、未读才显示待办。",
        "复核还确认了删除最后存档、返回标题和加载动画的边界行为，避免修复一个问题又引入新的流程问题。"
    ],
    "新手教程完整流程": [
        "新手教程采用模块式引导，不是单纯文字说明。被讲解的模块高亮，其余区域变暗；说明型步骤有下一步按钮，交互型步骤必须点击高亮区域才能继续。",
        "教程先在主页解释导航和开始新的一天，再进入游戏页演示顾客、订单、原料、当前饮品和制作按钮。教程营业是模拟流程，不计入正式收益。",
        "完成教程后弹出奖励框，发放初始资金和已解锁原料各 2 份。这个奖励和跳过教程保持一致，保证新玩家和熟练玩家起点公平。"
    ],
    "进入下一天 bug 修复": [
        "这个 bug 的本质是“日结已完成”和“下一天启动”之间的状态没有衔接好。前端进入下一天后页面到了 game.html，但 dayStarted 没有重新置为可启动状态。",
        "修复思路是把 startNextDay 明确写成状态重置函数：关闭结算弹窗、清除计时器、重置 dayStats、清除当前顾客和饮品，再调用 beginNewDay。",
        "同时后端 end-day 返回更新后的存档信息，前端重新读取玩家状态，保证天数、金币、经验和等级同步。"
    ],
    "进入下一天问题复测": [
        "用户重复反馈同一问题时，我没有只看单点代码，而是重新核对“已营业几天”和“局内第几天”的概念。",
        "主页显示的是 business_days，也就是已经完成结算的天数；局内显示的是 business_days + 1，也就是正在营业的第几天。",
        "复测后确认：第 0 天主页进入游戏显示第 1 天；完成结算后主页显示已营业 1 天；点击进入下一天后局内显示第 2 天。"
    ],
    "特殊道具 SQL bug": [
        "截图中的错误是 SQLITE no such column，说明 SQL 语句引用了不存在的列。定位后发现某处更新 save_collection 时误用了 monster_affinity.save_id。",
        "修复时把 WHERE 条件改回当前表真实字段，并顺带检查类似 SQL 更新，确保每张 save_* 表都用自己的 save_id 字段。",
        "这个问题说明后期表结构增多后，SQL 语句必须更谨慎。相同含义的 save_id 出现在不同表中，但不能跨表乱写列名。"
    ],
    "收尾、删除无用页面和日志页": [
        "收尾阶段先清理无用页面，避免导航里还有设置页、报告页等后续不用的入口。页面越少，答辩时越容易讲清楚主流程。",
        "AI 协作日志页作为课程展示功能保留。它不是游戏核心玩法，但能记录开发过程中的提示词、AI 摘要、人工修改和运行结果。",
        "日志接口提供查询、新增、删除三类操作。前端表单提交后刷新列表，删除时调用 DELETE 接口，形成一个完整的小型 CRUD 模块。"
    ],
    "整理全部提示词入日志": [
        "为了让日志页有真实展示内容，我把本次对话的阶段性需求整理成 23 条开发日志。每条日志包含提示词、AI 返回摘要、人工修改和运行结果。",
        "写入脚本需要兼容旧版 dev_logs 表。如果旧表里还存在 title、content、summary 等旧列，脚本会按实际列动态插入，避免 NOT NULL 约束报错。",
        "这一步的意义是把开发过程本身也变成系统数据，答辩时可以直接打开日志页说明项目是如何一步一步迭代完成的。"
    ],
    "修复日志乱码和时间显示": [
        "日志乱码不是前端样式问题，而是数据写入数据库前已经被错误编码。解决方式是用 UTF-8 脚本重新写入日志，而不是只在页面上替换字符。",
        "用户不想显示时间，因此后端 /api/logs 查询时不再返回 created_at，前端卡片也不渲染时间字段。",
        "修复后用真实接口验证，返回 23 条日志，并确认响应中没有 created_at 字段，页面也不再出现时间。"
    ],
    "实验报告初稿": [
        "报告初稿的任务是把项目从代码转化成可提交文档。因此我先运行本地服务，用 Playwright 打开关键页面截图，再用 python-docx 生成 Word。",
        "初稿覆盖了任务描述、系统设计、实验过程、功能截图、接口验证和心得，但每个步骤写得偏概括，适合作为框架，不足以满足详细课程报告。",
        "这一步新增了报告生成辅助脚本和截图目录，不改变游戏业务代码。"
    ],
    "实验报告详细化": [
        "本次扩写根据用户反馈，把报告从概要版升级成详细版。重点不再只是列出做了什么，而是说明为什么这样做、文件如何变化、状态如何流动、代码如何组织。",
        "截图部分不再只放图，而是给出每个界面的来源、点击路径、界面作用和关键交互。新手教程单独按步骤截图，便于展示引导流程。",
        "代码附录也从少量片段扩展为多组关键机制说明，包括 Express 托管、存档隔离、订单生成、评分、补货、日结、日志等。"
    ],
}


TUTORIAL_SHOTS = [
    ("tutorial_home_01_return_title.png", "主页教程 1：高亮“返回标题”，说明该按钮用于退出当前存档登录状态并返回标题页。"),
    ("tutorial_home_02_collection.png", "主页教程 2：高亮“怪兽图鉴”，说明图鉴记录来店怪兽和故事待办。"),
    ("tutorial_home_03_achievements.png", "主页教程 3：高亮“成就图鉴”，说明完成但未领取的成就会显示红点。"),
    ("tutorial_home_04_level_rewards.png", "主页教程 4：高亮“等级奖励”，说明升级后可领取阶段奖励。"),
    ("tutorial_home_05_backpack.png", "主页教程 5：高亮“查看背包”，说明物品、礼物和奖励会在背包集中查看。"),
    ("tutorial_home_06_shop.png", "主页教程 6：高亮“原料商店”，说明库存不足时可补货。"),
    ("tutorial_home_07_logs.png", "主页教程 7：高亮“开发日志”，说明该页记录 Codex 协作开发过程。"),
    ("tutorial_home_08_start_day.png", "主页教程 8：高亮“开始新的一天”，这是教程进入模拟营业的必点入口。"),
    ("tutorial_game_01_customer.png", "游戏教程 1：高亮顾客面板，说明怪兽名称、种族、心情和台词的位置。"),
    ("tutorial_game_02_order.png", "游戏教程 2：高亮订单需求，说明需求标签如何影响满意度。"),
    ("tutorial_game_03_ingredients.png", "游戏教程 3：高亮原料区，要求玩家选择带有目标口味标签的原料。"),
    ("tutorial_game_04_current_drink.png", "游戏教程 4：高亮当前饮品，说明已选原料会进入制作台。"),
    ("tutorial_game_05_make_button.png", "游戏教程 5：高亮制作按钮，点击后完成教程模拟饮品。"),
]


RESULT_SHOTS = [
    ("01_start.png", "标题页来自访问 start.html。玩家打开网站首先看到 Monster Café 艺术字标题，可以点击“开始游戏”进入存档页，或点击“退出游戏”弹出确认框。"),
    ("02_exit_game_modal.png", "点击标题页“退出游戏”后出现该弹窗，用于确认是否关闭网页。这个弹窗防止误触退出。"),
    ("02_saves.png", "点击标题页“开始游戏”进入 saves.html。这里展示所有存档，单击存档主体读取，底部加号条用于新建存档。"),
    ("03_new_save_modal.png", "点击存档页底部加号条后出现新建存档弹窗。玩家可输入昵称，也可点击骰子随机昵称。"),
    ("04_delete_save_modal.png", "鼠标悬停存档后右上角出现删除区域，点击垃圾桶进入确认弹窗。勾选本次登录不再提醒后，本次会话再次删除不再弹窗。"),
    ("03_home.png", "读取存档后进入 index.html 咖啡馆主页。主页显示玩家昵称、已营业天数、等级经验条、成就收集进度，并提供开始新一天、读取工作进度和商店入口。"),
    ("05_loading_overlay.png", "外部页面切换时显示加载过场。它由三枚布丁图标、进度条和随机提示语组成；营业中不会触发，避免影响计时。"),
    ("06_tutorial_choice_modal.png", "新建存档第一次进入主页时弹出新手教程选择框。进入或跳过教程都会发放相同初始奖励。"),
    ("08_return_title_modal.png", "在主页点击“返回标题”会弹出确认框。确认后清理当前存档登录状态，回到标题页。"),
    ("09_game_countdown.png", "点击“开始新的一天”进入 game.html 后先播放 3、2、1 倒计时，再正式开始顾客等待计时。"),
    ("04_game_customer.png", "正式营业界面来自 game.html。上方状态栏显示金币、第几天、等级、经验、剩余时间和今日顾客；中间是顾客、订单和评价墙；下方是可选原料。"),
    ("05_game_result.png", "玩家选择原料并点击“制作饮品”后出现结果。系统展示饮品名、满意度、金币奖励、经验和顾客评价，并把评价贴到评价墙。"),
    ("06_pause.png", "营业中点击“暂停”打开暂停弹窗。暂停会停止营业计时和顾客等待，玩家可继续营业、前往商店补货或退出本天。"),
    ("11_exit_day_confirm.png", "暂停弹窗中点击“退出本天”后出现二次确认。玩家可以保存进度并退出，也可以不保存直接回主页。"),
    ("12_day_settlement_modal.png", "一天结束后出现今日结算弹窗。它汇总服务顾客、差评、收入、成本、利润、平均满意度、经验和星级评价。"),
    ("07_shop.png", "从主页或补货入口进入 shop.html。商店展示原料名称、标签、价格、库存和解锁状态，玩家可购买或解锁原料。"),
    ("08_restock_shop.png", "从暂停弹窗选择补货进入商店时，页面进入补货状态。普通导航隐藏，只保留“补货完成继续营业”按钮。"),
    ("09_collection.png", "点击导航栏“怪兽图鉴”进入 collection.html。已解锁怪兽显示资料，未解锁怪兽显示问号卡。"),
    ("15_collection_detail.png", "点击某张已解锁怪兽卡片打开详情弹窗。这里展示种族、偏好、来店次数、最高满意度、好感和故事列表。"),
    ("16_story_reward_popup.png", "在怪兽详情中阅读未读故事后弹出奖励框，奖励会进入背包，图鉴待办随之消失。"),
    ("10_achievements.png", "点击“成就图鉴”进入 achievements.html。完成未领取的成就会排在前面，并在卡片和导航上显示待办。"),
    ("17_achievement_reward_popup.png", "点击可领取成就后弹出奖励框，显示本次领取到的金币、经验或特殊物品。"),
    ("11_level_rewards.png", "点击“等级奖励”进入 level-rewards.html。页面按等级展示奖励，当前等级达到且未领取时可点击领取。"),
    ("18_level_reward_popup.png", "点击可领取等级奖励后弹出奖励框。领取后该等级奖励会变为已领取并后置。"),
    ("12_backpack.png", "点击“查看背包”进入 backpack.html。左侧是物品图标和数量角标，右侧显示当前选中物品的名称、数量和简介。"),
    ("13_logs_form.png", "点击“开发日志”进入 logs.html。上方表单可新增提示词、AI 返回摘要、人工修改和运行结果。"),
    ("14_logs_list.png", "向下滚动日志页可看到日志列表。这里按对话阶段记录了 Codex 协作开发过程，不显示时间。"),
]


CODE_APPENDIX = [
    (
        "1. Express 启动与静态托管",
        "backend/server.js",
        dedent("""\
        const frontendDir = path.join(__dirname, '..', 'frontend');

        app.get('/', (req, res) => {
          res.sendFile(path.join(frontendDir, 'start.html'));
        });

        app.use('/api', apiRoutes);
        app.use(express.static(frontendDir, {
          etag: false,
          lastModified: false,
          maxAge: 0
        }));

        initDatabase()
          .then(() => startServer(PORT))
          .catch((error) => {
            console.error('数据库初始化失败：', error);
            process.exit(1);
          });
        """),
        "这段代码体现了项目的入口设计：根路径直接打开 start.html，/api 交给后端路由，其他静态文件由 Express 托管。启动服务前先执行 initDatabase，确保数据库表和默认数据准备完成。"
    ),
    (
        "2. SQLite 初始化与字段兼容",
        "backend/database.js",
        dedent("""\
        async function ensureColumn(tableName, columnName, columnSql) {
          const columns = await getColumnNames(tableName);
          if (!columns.includes(columnName)) {
            await run(`ALTER TABLE ${tableName} ADD COLUMN ${columnSql}`);
          }
        }

        async function initDatabase() {
          await createTables();
          await migrateTables();
          await seedInitialData();
        }
        """),
        "课程项目在迭代中不断加字段，如果每次都删除数据库会丢失演示数据。ensureColumn 用 PRAGMA table_info 检查字段是否存在，缺少时才 ALTER TABLE，保证旧数据库也能升级。"
    ),
    (
        "3. 前端请求自动绑定当前存档",
        "frontend/app.js",
        dedent("""\
        async function fetchJson(url, options = {}) {
          const selectedSave = localStorage.getItem('monsterCafeSelectedSave');
          const headers = { ...(options.headers || {}) };
          if (selectedSave) {
            headers['X-Save-Id'] = selectedSave;
          }
          const response = await fetch(url, { ...options, headers });
          if (!response.ok) throw new Error('请求失败');
          return response.json();
        }
        """),
        "多存档隔离依赖当前存档 ID。前端把 monsterCafeSelectedSave 存在 localStorage 中，每次请求都自动加到 X-Save-Id 请求头，后端据此读写对应存档的数据。"
    ),
    (
        "4. 存档维度初始化",
        "backend/routes/api.js",
        dedent("""\
        async function ensureSaveScopedData(saveId) {
          const save = await get('SELECT id FROM saves WHERE id = ?', [saveId]);
          if (!save) return false;

          const ingredients = await all('SELECT id, unlocked FROM ingredients');
          for (const ingredient of ingredients) {
            await run(
              `INSERT OR IGNORE INTO save_ingredients
               (save_id, ingredient_id, stock, unlocked)
               VALUES (?, ?, ?, ?)`,
              [saveId, ingredient.id, ingredient.unlocked ? 2 : 0, ingredient.unlocked ? 1 : 0]
            );
          }
          return true;
        }
        """),
        "这个函数是多存档架构的底座。新建或读取存档后，系统会为该存档准备库存、图鉴、成就、背包等奖励数据。INSERT OR IGNORE 避免重复初始化。"
    ),
    (
        "5. 随机顾客和订单生成",
        "backend/routes/api.js",
        dedent("""\
        router.get('/game/next-customer', async (req, res, next) => {
          const saveId = getRequestSaveId(req);
          await ensureSaveScopedData(saveId);
          const save = await get('SELECT id, level FROM saves WHERE id = ?', [saveId]);
          const monsters = await all(`
            SELECT monsters.*, save_collection.visit_count
            FROM monsters
            LEFT JOIN save_collection
              ON save_collection.monster_id = monsters.id
             AND save_collection.save_id = ?
            WHERE monsters.unlock_level <= ?
          `, [saveId, save.level]);
          const order = buildSolvableOrder(monster, unlockedIngredients);
          res.json({ monster, order_description: order.text, recommended_taste_tags: order.tags });
        });
        """),
        "顾客接口先按等级筛选可出现怪兽，再生成订单。订单不是随意拼接，而是基于已解锁原料判断是否可解，避免玩家遇到根本无法满分的需求。"
    ),
    (
        "6. 评分算法核心",
        "backend/routes/api.js",
        dedent("""\
        const selectedTags = new Set();
        for (const ingredient of selectedIngredients) {
          parseTasteTags(ingredient.taste_tags).forEach((tag) => selectedTags.add(tag));
        }

        const requiredHits = requestedTags.filter((tag) => selectedTags.has(tag)).length;
        const likedHits = likedTags.filter((tag) => selectedTags.has(tag)).length;
        const dislikedHits = dislikedTags.filter((tag) => selectedTags.has(tag)).length;

        let satisfaction = 35 + requiredHits * 25 + likedHits * 10 - dislikedHits * 20 + comboBonus;
        satisfaction = Math.max(0, Math.min(100, satisfaction));
        """),
        "评分先把所有原料标签合并成集合，因此第一标签和第二标签完全平等。需求命中是主要得分来源，喜欢标签提供额外加成，讨厌标签扣分，最后限制在 0 到 100。"
    ),
    (
        "7. 暂停与补货状态保存",
        "frontend/app.js",
        dedent("""\
        function saveGameState(mode = 'saved-day') {
          const state = {
            mode,
            remainingSeconds,
            customerWaitSeconds,
            dayStarted,
            dayStats: { ...dayStats },
            currentCustomer,
            currentOrderDescription,
            currentRequestedTasteTags,
            currentDrinkIngredientIds: currentDrinkIngredients.map((item) => item.id)
          };
          localStorage.setItem(getSaveStateKey(), JSON.stringify(state));
        }

        function goRestockFromGame() {
          pauseGame(false);
          saveGameState('restocking');
          window.location.href = 'shop.html?restock=1';
        }
        """),
        "补货不是结束本天，而是暂停后临时跳转。因此需要把局内状态保存到 localStorage。商店补货完成后，game.html 能恢复当前顾客、剩余时间和今日统计。"
    ),
    (
        "8. 日结计算与正式入账",
        "backend/routes/api.js",
        dedent("""\
        const profit = totalIncome - totalIngredientCost;
        const profitStars = getStars(profit, [100, 70, 40, 15]);
        const customerStars = getStars(reviewedCustomers, [8, 6, 4, 2]);
        const satisfactionStars = getStars(avgSatisfaction, [90, 80, 65, 50]);
        const totalXpGained = drinkXpTotal + settlementXp;

        await run(`
          UPDATE saves
          SET business_days = business_days + 1,
              coins = coins + ?,
              xp = xp + ?,
              level = ?,
              updated_at = CURRENT_TIMESTAMP
          WHERE id = ?
        `, [totalCoinGained, totalXpGained, nextLevel, saveId]);
        """),
        "正式收益在日结时统一写入 saves 表。这样保存退出、教程模式和局内临时显示不会提前污染存档数据。星级评分同时用于展示和经验奖励。"
    ),
    (
        "9. 成就排序和待办判断",
        "frontend/app.js",
        dedent("""\
        achievements.sort((a, b) => {
          if (a.completed && !a.claimed) return -1;
          if (b.completed && !b.claimed) return 1;
          if (!a.completed && b.claimed) return -1;
          if (!b.completed && a.claimed) return 1;
          return a.id - b.id;
        });

        bubble.hidden = !achievements.some((item) => item.completed && !item.claimed);
        """),
        "待办红点只对“已完成但未领取”的内容出现。未完成成就不应该让导航栏标红，否则新档一进来就会出现错误引导。排序也按可领取、未完成、已领取分层。"
    ),
    (
        "10. 怪兽故事奖励",
        "backend/routes/api.js",
        dedent("""\
        router.post('/collection/:monsterId/read-story', async (req, res, next) => {
          const saveId = getRequestSaveId(req);
          const threshold = Number(req.body.threshold);
          const row = await get(`
            SELECT monster_affinity.read_stories, monster_affinity.affinity
            FROM monster_affinity
            WHERE save_id = ? AND monster_id = ?
          `, [saveId, monsterId]);
          const rewards = await addBackpackReward(saveId, reward.code, reward.name, reward.description, reward.icon_text, 1);
          res.json({ ok: true, rewards: [rewards] });
        });
        """),
        "故事奖励需要同时更新 read_stories 和背包。只有好感达到阈值且故事未读时才允许领取，领取后奖励物品进入 save_backpack，图鉴红点消失。"
    ),
    (
        "11. 开发日志接口",
        "backend/routes/api.js",
        dedent("""\
        router.get('/logs', async (req, res, next) => {
          const logs = await all(`
            SELECT id, prompt, ai_summary, manual_change, run_result
            FROM dev_logs
            ORDER BY id ASC
          `);
          res.json(logs);
        });
        """),
        "日志接口只返回展示需要的四个字段，不返回 created_at。这样页面不会显示时间，也避免用户看到不需要的数据库字段。"
    ),
]


def set_font(run, east="宋体", size=12, bold=False, color=None):
    run.font.name = east
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), east)
    r_fonts.set(qn("w:ascii"), "Calibri" if east != "Consolas" else "Consolas")
    r_fonts.set(qn("w:hAnsi"), "Calibri" if east != "Consolas" else "Consolas")


def para(doc, text="", size=12, bold=False, color=None, align=None, before=0, after=6):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, east="黑体", size={1: 16, 2: 14, 3: 12.5}.get(level, 12), bold=True, color="2E2143")
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r, size=12)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    return p


def image(doc, filename, caption, width=5.7):
    path = SCREENSHOT_DIR / filename
    if not path.exists():
        para(doc, f"（缺少截图：{filename}）", color="9B1C1C")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width))
    para(doc, caption, size=10.5, color="6B4C72", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        shade_cell(t.rows[0].cells[i], "FCE8F1")
        cell_text(t.rows[0].cells[i], h, bold=True, size=size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cell_text(cells[i], v, size=size)
    para(doc, "", after=2)
    return t


def code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(4)
    for line in code.splitlines():
        r = p.add_run(line + "\n")
        set_font(r, east="Consolas", size=8.5)
    return p


def write_process_step(doc, index, step):
    heading(doc, f"{index}. {step['title']}", 2)
    para(doc, f"精简提示词：{step['prompt']}", bold=True, color="2E2143")
    para(doc, f"本阶段 Codex 的工作目标：{step['work']}")
    para(doc, f"涉及文件：{step['files']}")
    para(doc, f"架构变化：{step['architecture']}")
    para(doc, "工作思路与实现逻辑：", bold=True, color="2E2143", after=3)
    for note in DETAIL_NOTES.get(step["title"], [step["logic"]]):
        bullet(doc, note)
    para(doc, "代码逻辑说明：", bold=True, color="2E2143", after=3)
    para(
        doc,
        f"这一部分的代码重点可以概括为：{step['logic']}。实现时没有把逻辑写死在页面文字里，而是尽量让后端负责数据判断和状态持久化，前端负责事件绑定、界面刷新和用户反馈。这样做的好处是后续需求变化时，只需要在对应层修改规则或展示，不会让所有页面互相耦合。"
    )
    para(
        doc,
        "测试和调整方式：开发完成后通过本地 npm start 启动服务，在浏览器中直接点击页面流程，并结合接口返回、页面截图和数据库状态判断是否符合需求。发现问题后优先判断是数据错误、接口错误还是前端状态错误，再做定点修复。"
    )
    doc.add_page_break()


def screenshot_explanation(doc, filename, caption):
    heading(doc, caption.split("：")[0], 3)
    image(doc, filename, caption, width=5.2)
    para(doc, f"来源与入口：该界面由本地服务运行后截图，文件为 {filename}。用户需要按照界面对应入口点击，例如标题页、存档页、主页导航、暂停弹窗或奖励卡片，才能进入该状态。")
    para(doc, f"界面作用：{caption} 这张图不是单纯装饰截图，而是用来说明功能是否已经接入实际页面、交互按钮是否可见、中文文案是否正常，以及像素风卡片布局是否统一。")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Monster Café 怪兽咖啡馆经营小游戏扩写实验报告")
    set_font(r, east="黑体", size=21, bold=True, color="2E2143")
    para(doc, "基于 HTML、CSS、JavaScript、Node.js、Express 与 SQLite 的本地网页小游戏系统", align=WD_ALIGN_PARAGRAPH.CENTER, color="6B4C72")

    table(doc, ["项目项", "内容"], [
        ["项目名称", "Monster Café 怪兽咖啡馆经营小游戏"],
        ["开发模式", "用户提出需求与测试反馈，Codex 作为全栈开发智能体辅助实现、调试、截图和报告整理"],
        ["技术栈", "前端 HTML/CSS/JavaScript；后端 Node.js + Express；数据库 SQLite"],
        ["报告重点", "按提示词顺序扩写工作思路、代码逻辑、文件变化、架构变化和最终功能展示"],
    ])

    heading(doc, "一、智能体辅助说明", 1)
    para(doc, "本项目使用 Codex 作为全栈开发智能体辅助。我的工作方式不是直接一次性生成所有代码，而是根据用户每次提出的阶段性需求，先判断该需求影响的是数据结构、后端接口、前端页面、交互状态还是视觉样式，然后再进行对应修改。")
    para(doc, "在整个开发过程中，用户负责提出课程设计目标、界面偏好、功能要求和测试反馈；Codex 负责读取现有项目、拆解需求、修改文件、运行服务、定位错误、生成截图和整理报告。每次修复都尽量沿用已有结构，避免无关重构。")
    para(doc, "这种协作方式的特点是迭代清晰：从项目骨架开始，逐步加入数据库、API、游戏页、商店、图鉴、成就、背包、新手教程、日志和报告。每一步都能独立运行和验证，最终组合成完整系统。")

    heading(doc, "二、任务描述", 1)
    para(doc, "Monster Café 是一个怪兽咖啡馆经营小游戏。玩家扮演咖啡馆老板，每天接待不同的小怪兽顾客。顾客提出带有口味偏好的饮品需求，玩家选择原料调配饮品，系统根据原料标签与怪兽偏好计算满意度，并给予金币、经验、图鉴、成就或物品奖励。")
    para(doc, "项目要求本地运行，前端不使用复杂框架，后端使用 Node.js + Express，数据库使用 SQLite。页面风格偏可爱、像素风、卡片式布局，所有实际交互文字使用中文，Monster Café 和怪兽名等特殊名称保留。")
    bullet(doc, "核心玩法：生成顾客、查看订单、选择原料、制作饮品、获得评价。")
    bullet(doc, "经营系统：库存、商店、补货、日结、金币、经验、等级。")
    bullet(doc, "收集系统：怪兽图鉴、故事、成就图鉴、等级奖励、背包。")
    bullet(doc, "展示系统：AI 协作开发日志和实验报告截图。")

    heading(doc, "三、最初代码架构与后续演变", 1)
    para(doc, "项目最初采用 backend 与 frontend 分离的目录结构。backend 负责 Express 服务、SQLite 初始化和 API；frontend 负责 HTML 页面、CSS 样式和原生 JavaScript 交互。这个结构简单清晰，适合课程设计展示，也方便后续逐步扩展。")
    table(doc, ["阶段", "架构状态", "变化说明"], [
        ["初始骨架", "backend/server.js + database.js + routes；frontend 多个 HTML 页面", "先形成本地可运行工程"],
        ["基础游戏", "新增游戏 API 与 game.html 交互", "从静态页面变成可玩流程"],
        ["经营扩展", "加入商店、日结、成就、更多数据表", "从单次调饮扩展为按天经营"],
        ["多存档重构", "引入 saves 和 save_* 系列表", "金币、库存、图鉴、成就、背包按存档隔离"],
        ["完整应用", "加入标题页、存档页、主页、教程、奖励、日志", "从小游戏扩展为完整课程展示系统"],
    ])

    doc.add_page_break()
    heading(doc, "四、按提示词顺序扩写的实验过程", 1)
    para(doc, "以下内容按照对话提示词顺序展开。每一小节都说明该阶段用户提出了什么、Codex 如何判断工作重点、具体改动了哪些文件、架构是否变化，以及代码逻辑如何实现。")
    doc.add_page_break()
    for i, step in enumerate(PROCESS_STEPS, start=1):
        write_process_step(doc, i, step)

    heading(doc, "五、新手教程逐步截图与讲解", 1)
    para(doc, "新手教程是本项目后期补充的重要交互。它不是普通说明文字，而是通过高亮目标模块、遮罩其他区域、显示对话框的方式引导玩家操作。说明型步骤可以点击“下一步”，交互型步骤必须点击高亮模块才能继续。")
    for filename, caption in TUTORIAL_SHOTS:
        screenshot_explanation(doc, filename, caption)

    doc.add_page_break()
    heading(doc, "六、最终运行结果截图与讲解", 1)
    para(doc, "本节展示最终系统的主要功能页面。每张截图都说明它来自哪个入口、用户如何触发，以及该界面承担的功能。")
    for filename, caption in RESULT_SHOTS:
        screenshot_explanation(doc, filename, caption)

    heading(doc, "七、关键代码附录与讲解", 1)
    para(doc, "附录选取项目中最能体现实现机制的代码片段。每段代码后都附有解释，说明它在系统中承担的职责以及为什么这样写。附录控制在 20 页以内，重点讲核心机制，不贴完整文件。")
    for title, path, code, explanation in CODE_APPENDIX:
        heading(doc, title, 2)
        para(doc, f"文件位置：{path}", color="6B4C72", size=10.5)
        code_block(doc, code.strip())
        para(doc, f"代码讲解：{explanation}")

    heading(doc, "八、收获与心得", 1)
    para(doc, "这次实验让我更清楚地认识到，一个课程设计项目不能只满足“能打开页面”或“能调用接口”。真正完整的项目需要清晰的数据模型、可解释的游戏规则、稳定的状态管理、友好的交互提示，以及可以展示开发过程的文档。")
    para(doc, "开发中最重要的一次架构变化是多存档隔离。最初很多数据都可以看作全局数据，但只要有多个存档，金币、经验、库存、图鉴、成就和背包都必须独立，否则新档会继承旧档数据，游戏逻辑就会失真。")
    para(doc, "评分机制的调整也很有代表性。用户测试发现第二标签命中没有得到合理分数、某些需求可能无解，这说明游戏算法不仅要能算，还要符合玩家直觉。最终通过多标签平等、订单可解性判断和更合理的加减分，让系统更公平。")
    para(doc, "AI 协作的价值在于加快实现、定位问题和整理材料，但最终判断仍来自用户反馈。每次截图反馈都推动了一次具体修复，例如结算错误、补货返回失效、SQL 字段错误、日志乱码和报告内容过简等。")

    heading(doc, "九、本地运行方式", 1)
    for item in [
        "进入项目目录：C:\\Users\\xersn\\Documents\\vibe。",
        "执行 npm install 安装依赖。",
        "执行 npm start 启动 Express 服务。",
        "浏览器访问 http://localhost:3000，进入 Monster Café 标题页。",
        "选择或新建存档后，可以进入主页、开始营业、补货、查看图鉴、领取奖励、打开背包和记录开发日志。",
    ]:
        bullet(doc, item)

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
